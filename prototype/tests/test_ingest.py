"""Clause-aware chunker (breadcrumbs, clause numbers, pages) and regex metadata."""

import io
from datetime import date

from mixup.ingest import chunk_document, extract_metadata, parse_bytes, parse_file
from mixup.models import Document

CIRCULAR = """SINTETIK - CONTOH SAHAJA / SYNTHETIC - SAMPLE ONLY
SURAT PEKELILING PERKHIDMATAN BILANGAN 7 TAHUN 2024
Rujukan: SPP 7/2024
Tarikh kuat kuasa: 1 Mac 2024

TUJUAN
1. Surat pekeliling ini bertujuan menerangkan peraturan contoh untuk ujian unit sahaja.

PEMAKAIAN
2. Pekeliling ini terpakai kepada semua pegawai Perkhidmatan Awam Persekutuan.
\f
PERATURAN
3. Syarat-syarat adalah seperti berikut:
3.1 Pegawai hendaklah memohon dalam tempoh 14 hari sebelum tarikh berkenaan.
(a) permohonan dibuat secara dalam talian;
(b) disokong oleh penyelia.
3.2 Kelulusan diberi oleh Ketua Jabatan dalam tempoh 7 hari bekerja.
"""


def _doc():
    return Document(doc_id="T-1", circular_no="SPP 7/2024")


def test_plain_text_headings_paragraphs_and_pages():
    pages = parse_bytes("t.txt", CIRCULAR.encode(), "T-1")
    assert [p.page_no for p in pages] == [1, 2]
    chunks = chunk_document(_doc(), pages)
    by_clause = {c.clause_ref: c for c in chunks}
    assert by_clause["2"].breadcrumb == "SPP 7/2024 > PEMAKAIAN > 2"
    assert by_clause["2"].page_start == 1
    assert by_clause["3.1"].section == "PERATURAN" and by_clause["3.1"].page_start == 2
    # (a)/(b) items stay inside their paragraph
    assert "(b) disokong oleh penyelia" in by_clause["3.1"].text
    assert "3.2" in by_clause and "(a)" not in by_clause
    assert all("SINTETIK" not in c.text for c in chunks)


def test_short_lead_in_merged_into_next_clause():
    chunks = chunk_document(_doc(), parse_bytes("t.txt", CIRCULAR.encode(), "T-1"))
    clause_31 = next(c for c in chunks if c.clause_ref == "3.1")
    assert clause_31.text.startswith("3. Syarat-syarat")


def test_markdown_corpus_document(store):
    chunks = store.chunks_by_doc["SPP-1-2023"]
    by_clause = {c.clause_ref: c for c in chunks}
    assert by_clause["3"].breadcrumb == "SPP 1/2023 > PEMAKAIAN > 3"
    assert by_clause["4.1"].breadcrumb == "SPP 1/2023 > PERATURAN TUNTUTAN > 4.1"
    assert "60 hari" in by_clause["4.1"].text and by_clause["4.1"].page_start == 5
    assert "RM0.70" in by_clause["4.2"].text
    assert by_clause["6"].section == "PEMBATALAN" and by_clause["6"].page_start == 7
    assert chunks[0].clause_ref == "Tajuk"


def test_every_document_is_watermarked(store):
    for doc_id, pages in store.pages.items():
        lines = [x for x in "\n".join(p.text for p in pages).split("\n") if x.strip()]
        assert lines[0].startswith("SINTETIK - CONTOH SAHAJA"), doc_id
        assert lines[-1].startswith("SINTETIK - CONTOH SAHAJA"), doc_id


def test_metadata_regex_on_upload_demo(store):
    pages = parse_file(store.data_dir / "upload_demo" / "SPP-1-2026.md")
    meta = extract_metadata(pages)
    assert meta["circular_no"] == "SPP 1/2026"
    assert meta["series"] == "SPP"
    assert meta["effective_date"] == date(2026, 10, 1)
    assert meta["issue_date"] == date(2026, 9, 15)
    assert meta["jurisdiction"] == "FEDERAL"
    assert meta["cluster"] == "travel-claims"
    assert meta["classification_level"] == 0
    assert meta["language"] == "ms"


def test_metadata_jurisdiction_and_classification(store):
    def meta(doc_id):
        return extract_metadata(store.pages[doc_id])

    assert meta("PAN-2-2024")["jurisdiction"] == "SARAWAK"
    assert meta("PP-3-2024")["jurisdiction"] == "FEDERAL_SARAWAK"
    assert meta("SOP-AUDIT-2025")["classification_level"] == 1
    assert meta("PKP-1-2025")["classification_level"] == 2
    assert meta("SE-1-2021").get("one_off") is True
    assert meta("GP-2-2025")["language"] == "en"


def test_manual_metadata_wins(store):
    # metadata.csv says PKP 1/2025 is an SOP; the regex guess differs but the manual value wins
    assert store.docs["PKP-1-2025"].doc_type == "sop"
    assert store.docs["PAN-2-2024"].jurisdiction == "SARAWAK"


def test_docx_parsing():
    import docx

    d = docx.Document()
    d.add_heading("SURAT PEKELILING PERKHIDMATAN BILANGAN 9 TAHUN 2024", level=1)
    d.add_heading("PEMAKAIAN", level=2)
    d.add_paragraph("1. Pekeliling ini terpakai kepada semua pegawai Perkhidmatan Awam Persekutuan.")
    buf = io.BytesIO()
    d.save(buf)
    pages = parse_bytes("x.docx", buf.getvalue(), "X")
    chunks = chunk_document(Document(doc_id="X", circular_no="SPP 9/2024"), pages)
    assert any(c.breadcrumb == "SPP 9/2024 > PEMAKAIAN > 1" for c in chunks)


def test_blank_pdf_page_flagged_for_ocr():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    pages = parse_bytes("scan.pdf", buf.getvalue(), "SCAN")
    assert len(pages) == 1 and pages[0].needs_ocr
